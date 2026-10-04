"""
Recurring Report Automation (Weekly Performance + Monthly Advertiser Reconciliation)
Project: AP-TRIS

Replaces two manual Excel routines of an affiliate operations team:

1. WEEKLY PERFORMANCE REPORT  (every Monday)
   KPI per offer for the last complete week vs. the previous week (WoW),
   with automatic flags for campaigns that need attention.

2. MONTHLY ADVERTISER RECONCILIATION  (first working days of each month)
   Matches the platform's tracking data with each advertiser's own report
   (joined on click_id = sub_id), classifies every discrepancy, computes the
   billable amount per advertiser and the publisher payout statement
   (based on advertiser-confirmed conversions, PIT 10% withholding, fraud hold).

Usage:
    python 05_pipeline_automation/recurring_reports.py                 # latest week + latest month
    python 05_pipeline_automation/recurring_reports.py --month 2026-08
    python 05_pipeline_automation/recurring_reports.py --week 2026-09-14  # week starting Monday 14/09

Schedule (examples):
    Windows Task Scheduler : run every Monday 07:00
    cron                   : 0 7 * * 1  python /path/recurring_reports.py

Outputs (Excel, formatted for business users):
    01_data/exports/reports/weekly_performance_<week_start>.xlsx
    01_data/exports/reports/monthly_reconciliation_<YYYY-MM>.xlsx
"""
import argparse
import os

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LAKE = os.path.join(BASE, "01_data", "lakehouse")
ADV_DIR = os.path.join(BASE, "01_data", "sources", "advertiser_reports")
OUT_DIR = os.path.join(BASE, "01_data", "exports", "reports")
VN_OFFSET = pd.Timedelta(hours=7)

FLAGGED_STATUS = "Flagged"
PIT_THRESHOLD_VND = 2_000_000        # withhold 10% PIT when monthly payout >= 2M VND
PIT_RATE = 0.10
WOW_AR_DROP_PP = 5.0                 # flag: approval rate drops > 5 percentage points
WOW_MARGIN_DROP_PCT = -20.0          # flag: margin drops > 20% WoW
BOT_LEADS_ALERT = 5                  # publisher watchlist: >= 5 leads converted in < 5 seconds in the week
LOW_AR_ALERT_PCT = 15.0              # publisher watchlist: approval rate < 15% with >= 20 leads
RECON_TOLERANCE_PCT = 3.0            # unexplained variance above 3% is escalated to the advertiser


# ---------------------------------------------------------------- loading
def load():
    conv = pd.read_csv(os.path.join(LAKE, "silver", "fact_conversions_cleansed.csv.gz"))
    clicks = pd.read_csv(os.path.join(LAKE, "bronze", "fact_clicks.csv.gz"),
                         usecols=["click_time", "offer_id"])
    offers = pd.read_csv(os.path.join(LAKE, "silver", "dim_offers.csv.gz"))
    pubs = pd.read_csv(os.path.join(LAKE, "silver", "dim_publishers.csv.gz"))
    conv["conversion_time"] = pd.to_datetime(conv["conversion_time"]) + VN_OFFSET
    clicks["click_time"] = pd.to_datetime(clicks["click_time"]) + VN_OFFSET
    return conv, clicks, offers, pubs


# ---------------------------------------------------------------- weekly
def weekly_kpi(conv, clicks, offers, start, end):
    c = conv[(conv.conversion_time >= start) & (conv.conversion_time < end)]
    k = clicks[(clicks.click_time >= start) & (clicks.click_time < end)]
    g = c.groupby("offer_id").agg(
        conversions=("conversion_id", "count"),
        approved=("status", lambda s: (s == "Approved").sum()),
        bot_leads=("time_to_convert_seconds", lambda s: (s < 5).sum()),
        revenue_vnd=("advertiser_revenue_vnd", "sum"),
        payout_vnd=("publisher_payout_vnd", "sum"),
        margin_vnd=("gross_margin_vnd", "sum"),
    )
    g["clicks"] = k.groupby("offer_id").size()
    g = offers.set_index("offer_id")[["offer_name", "vertical", "payout_model"]].join(g).fillna(0)
    g["cr_pct"] = np.where(g.clicks > 0, g.conversions / g.clicks * 100, 0)
    g["approval_rate_pct"] = np.where(g.conversions > 0, g.approved / g.conversions * 100, 0)
    g["epc_vnd"] = np.where(g.clicks > 0, g.payout_vnd / g.clicks, 0)
    return g


def build_weekly(conv, clicks, offers, week_start):
    ws = pd.Timestamp(week_start)
    cur = weekly_kpi(conv, clicks, offers, ws, ws + pd.Timedelta(days=7))
    prev = weekly_kpi(conv, clicks, offers, ws - pd.Timedelta(days=7), ws)

    df = cur.copy()
    df["margin_prev_vnd"] = prev["margin_vnd"]
    df["margin_wow_pct"] = np.where(prev.margin_vnd > 0, (cur.margin_vnd / prev.margin_vnd - 1) * 100, np.nan)
    df["approval_rate_prev_pct"] = prev["approval_rate_pct"]
    df["approval_rate_wow_pp"] = cur.approval_rate_pct - prev.approval_rate_pct
    df["bot_share_pct"] = np.where(cur.conversions > 0, cur.bot_leads / cur.conversions * 100, 0)

    def flag(r):
        f = []
        if r.approval_rate_wow_pp <= -WOW_AR_DROP_PP:
            f.append(f"Approval rate -{abs(r.approval_rate_wow_pp):.1f}pp")
        if pd.notna(r.margin_wow_pct) and r.margin_wow_pct <= WOW_MARGIN_DROP_PCT:
            f.append(f"Margin {r.margin_wow_pct:.0f}% WoW")
        return "; ".join(f) if f else "OK"

    df["flag"] = df.apply(flag, axis=1)
    df = df.reset_index().sort_values("margin_vnd", ascending=False)

    tot_c, tot_p = cur.sum(numeric_only=True), prev.sum(numeric_only=True)
    def pct(a, b): return (a / b - 1) * 100 if b else np.nan
    summary = pd.DataFrame([
        ["Clicks", tot_c.clicks, tot_p.clicks, pct(tot_c.clicks, tot_p.clicks)],
        ["Conversions", tot_c.conversions, tot_p.conversions, pct(tot_c.conversions, tot_p.conversions)],
        ["Approved conversions", tot_c.approved, tot_p.approved, pct(tot_c.approved, tot_p.approved)],
        ["Approval rate (%)", tot_c.approved / tot_c.conversions * 100, tot_p.approved / tot_p.conversions * 100, np.nan],
        ["Revenue (VND)", tot_c.revenue_vnd, tot_p.revenue_vnd, pct(tot_c.revenue_vnd, tot_p.revenue_vnd)],
        ["Gross margin (VND)", tot_c.margin_vnd, tot_p.margin_vnd, pct(tot_c.margin_vnd, tot_p.margin_vnd)],
        ["EPC (VND/click)", tot_c.payout_vnd / tot_c.clicks, tot_p.payout_vnd / tot_p.clicks, np.nan],
        ["Offers flagged", (df.flag != "OK").sum(), np.nan, np.nan],
    ], columns=["KPI", "This week", "Previous week", "WoW %"])

    wk = conv[(conv.conversion_time >= ws) & (conv.conversion_time < ws + pd.Timedelta(days=7))]
    pw = wk.groupby("publisher_id").agg(
        leads=("conversion_id", "count"),
        approved=("status", lambda s: (s == "Approved").sum()),
        bot_leads=("time_to_convert_seconds", lambda s: (s < 5).sum()),
        payout_vnd=("publisher_payout_vnd", "sum"))
    pw["approval_rate_pct"] = pw.approved / pw.leads * 100
    pw["reason"] = np.where(pw.bot_leads >= BOT_LEADS_ALERT, "BOT: leads converted in < 5s",
                   np.where((pw.leads >= 20) & (pw.approval_rate_pct < LOW_AR_ALERT_PCT), "LOW APPROVAL", ""))
    watch = pw[pw.reason != ""].reset_index().sort_values("bot_leads", ascending=False)
    summary.loc[len(summary)] = ["Publishers on watchlist", len(watch), np.nan, np.nan]

    cols = ["offer_id", "offer_name", "vertical", "payout_model", "clicks", "conversions", "cr_pct", "approved",
            "approval_rate_pct", "approval_rate_prev_pct", "approval_rate_wow_pp", "revenue_vnd", "margin_vnd",
            "margin_prev_vnd", "margin_wow_pct", "epc_vnd", "bot_share_pct", "flag"]
    return summary, df[cols], watch


# ---------------------------------------------------------------- monthly reconciliation
def build_recon(conv, offers, pubs, month):
    path = os.path.join(ADV_DIR, f"advertiser_report_{month}.csv.gz")
    adv = pd.read_csv(path)
    plat = conv[conv.conversion_time.dt.strftime("%Y-%m") == month].copy()
    plat["plat_status"] = plat["status"].replace({"Fraud": "Rejected"})
    plat["plat_revenue_vnd"] = plat["offer_id"].map(offers.set_index("offer_id")["advertiser_revenue_vnd"]) \
        .where(plat["plat_status"] == "Approved", 0.0)

    m = plat[["click_id", "conversion_id", "publisher_id", "offer_id", "advertiser_name", "conversion_time",
              "plat_status", "plat_revenue_vnd"]].merge(
        adv.rename(columns={"sub_id": "click_id", "offer_id": "adv_offer_id", "advertiser_name": "adv_name"}),
        on="click_id", how="outer", indicator=True)
    m["advertiser_name"] = m["advertiser_name"].fillna(m["adv_name"])
    m["offer_id"] = m["offer_id"].fillna(m["adv_offer_id"])
    m["adv_revenue_vnd"] = m["adv_revenue_vnd"].fillna(0.0)
    m["plat_revenue_vnd"] = m["plat_revenue_vnd"].fillna(0.0)

    def classify(r):
        if r["_merge"] == "left_only":
            return "MISSING_AT_ADVERTISER"
        if r["_merge"] == "right_only":
            return "MISSING_IN_TRACKING"
        if r.plat_status == "Pending":
            return "PENDING_FINALISED"
        if r.plat_status != r.adv_status:
            return "STATUS_MISMATCH"
        if abs(r.plat_revenue_vnd - r.adv_revenue_vnd) > 0.5:
            return "AMOUNT_MISMATCH"
        return "MATCHED"

    m["recon_status"] = m.apply(classify, axis=1)
    m["variance_vnd"] = m["adv_revenue_vnd"] - m["plat_revenue_vnd"]
    m.loc[m.recon_status == "PENDING_FINALISED", "variance_vnd"] = 0.0

    # summary per advertiser
    s = m.groupby("advertiser_name").agg(
        platform_leads=("conversion_id", "count"),
        advertiser_leads=("adv_status", "count"),
        platform_approved=("plat_status", lambda x: (x == "Approved").sum()),
        advertiser_approved=("adv_status", lambda x: (x == "Approved").sum()),
        platform_revenue_vnd=("plat_revenue_vnd", "sum"),
        billable_revenue_vnd=("adv_revenue_vnd", "sum"),
        matched=("recon_status", lambda x: (x == "MATCHED").sum()),
        discrepancies=("recon_status", lambda x: x.isin(["MISSING_AT_ADVERTISER", "MISSING_IN_TRACKING",
                                                           "STATUS_MISMATCH", "AMOUNT_MISMATCH"]).sum()),
    )
    pend = m[m.recon_status == "PENDING_FINALISED"].groupby("advertiser_name")["adv_revenue_vnd"].sum()
    s["pending_finalised_vnd"] = pend.reindex(s.index).fillna(0)
    # variance that cannot be explained by pending leads being finalised -> must be checked with advertiser
    s["variance_vnd"] = s.billable_revenue_vnd - s.platform_revenue_vnd - s.pending_finalised_vnd
    s["variance_pct"] = np.where(s.platform_revenue_vnd > 0, s.variance_vnd / s.platform_revenue_vnd * 100, 0)
    s["match_rate_pct"] = s.matched / (s.platform_leads.clip(lower=1)) * 100
    s["action"] = np.where(s.variance_pct.abs() > RECON_TOLERANCE_PCT, "ESCALATE - send discrepancy file",
                           "OK - issue invoice")
    s = s.reset_index().sort_values("billable_revenue_vnd", ascending=False)

    by_type = m.groupby("recon_status").agg(records=("click_id", "count"), variance_vnd=("variance_vnd", "sum")) \
        .reset_index().sort_values("records", ascending=False)

    disc = m[~m.recon_status.isin(["MATCHED", "PENDING_FINALISED"])].copy()
    disc = disc[["recon_status", "advertiser_name", "offer_id", "click_id", "conversion_id", "publisher_id",
                 "conversion_time", "plat_status", "adv_status", "plat_revenue_vnd", "adv_revenue_vnd", "variance_vnd"]] \
        .sort_values(["recon_status", "advertiser_name"])

    # publisher payout statement: pay only advertiser-confirmed approvals, hold flagged publishers
    payout_rate = offers.set_index("offer_id")["publisher_payout_vnd"]
    ok = m[(m.adv_status == "Approved") & m.publisher_id.notna()].copy()
    ok["payout_vnd"] = ok["offer_id"].map(payout_rate)
    p = ok.groupby("publisher_id").agg(confirmed_conversions=("click_id", "count"), gross_payout_vnd=("payout_vnd", "sum"))
    p = pubs.set_index("publisher_id")[["traffic_channel", "tier", "status"]].join(p, how="left").fillna(
        {"confirmed_conversions": 0, "gross_payout_vnd": 0})
    p["pit_withholding_vnd"] = np.where(p.gross_payout_vnd >= PIT_THRESHOLD_VND, p.gross_payout_vnd * PIT_RATE, 0)
    p["net_payable_vnd"] = p.gross_payout_vnd - p.pit_withholding_vnd
    p["payment_status"] = np.where(p.status == FLAGGED_STATUS, "HOLD - fraud audit", "READY TO PAY")
    p = p.reset_index().sort_values("gross_payout_vnd", ascending=False)

    tot = s[["platform_revenue_vnd", "pending_finalised_vnd", "billable_revenue_vnd", "variance_vnd"]].sum()
    overview = pd.DataFrame([
        ["Reconciliation month", month],
        ["Platform leads (tracking)", int(plat.shape[0])],
        ["Advertiser leads (reports)", int(adv.shape[0])],
        ["Matched records", int((m.recon_status == "MATCHED").sum())],
        ["Discrepancy records", int(s.discrepancies.sum())],
        ["Platform revenue (VND)", tot.platform_revenue_vnd],
        ["+ Pending leads finalised as approved (VND)", tot.pending_finalised_vnd],
        ["Billable revenue - advertiser confirmed (VND)", tot.billable_revenue_vnd],
        ["Unexplained variance (VND)", tot.variance_vnd],
        ["Unexplained variance (%)", tot.variance_vnd / tot.platform_revenue_vnd * 100],
        ["Advertisers to escalate (variance > 3%)", int((s.action.str.startswith("ESCALATE")).sum())],
        ["Publisher payout - ready to pay (VND)", p.loc[p.payment_status == "READY TO PAY", "net_payable_vnd"].sum()],
        ["Publisher payout - on hold (VND)", p.loc[p.payment_status != "READY TO PAY", "gross_payout_vnd"].sum()],
    ], columns=["Item", "Value"])
    return overview, s, by_type, disc, p


# ---------------------------------------------------------------- excel formatting
NAVY = "0F172A"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
RED_FILL = PatternFill("solid", fgColor="FDE2E2")
GREEN_FILL = PatternFill("solid", fgColor="DCFCE7")


def write_excel(path, sheets, title):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False, startrow=2)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws["A1"] = f"{title} - {ws.title}"
        ws["A1"].font = Font(bold=True, size=13, color=NAVY)
        for cell in ws[3]:
            cell.fill = HEADER_FILL
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[3].height = 32
        ws.freeze_panes = "A4"
        if ws.max_row > 3:
            ws.auto_filter.ref = f"A3:{get_column_letter(ws.max_column)}{ws.max_row}"
        for col in ws.iter_cols(min_row=3):
            header = str(col[0].value or "")
            width = max([len(header)] + [len(str(c.value)) for c in col[1:50] if c.value is not None]) + 2
            ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 42)
            for c in col[1:]:
                if isinstance(c.value, (int, float)):
                    if header.endswith("_pct") or header.endswith("_pp") or "%" in header:
                        c.number_format = "0.0"
                    elif header.endswith("_vnd") or "VND" in header or isinstance(c.value, float):
                        c.number_format = "#,##0"
                    else:
                        c.number_format = "#,##0"
        # highlight status columns
        for col in ws.iter_cols(min_row=3):
            h = str(col[0].value)
            if h in ("flag", "action", "payment_status", "recon_status"):
                for c in col[1:]:
                    v = str(c.value or "")
                    if v.startswith(("OK", "READY", "MATCHED")):
                        c.fill = GREEN_FILL
                    elif v:
                        c.fill = RED_FILL
    wb.save(path)


# ---------------------------------------------------------------- chart for README
def plot_recon(by_adv, by_type, month, path):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return
    labels = {"MISSING_AT_ADVERTISER": "Missing at advertiser", "STATUS_MISMATCH": "Status mismatch",
              "AMOUNT_MISMATCH": "Amount mismatch", "MISSING_IN_TRACKING": "Missing in tracking"}
    t = by_type[by_type.recon_status.isin(labels)].set_index("recon_status").reindex(list(labels)[::-1])
    a = by_adv.sort_values("variance_pct")
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 6.2), gridspec_kw={"width_ratios": [1, 1.4]})
    ax1.barh([labels[i] for i in t.index], t.records, color="#2563EB", height=0.55)
    for i, (n, v) in enumerate(zip(t.records, t.variance_vnd)):
        ax1.text(n + 8, i, f"{n:,} records | {v / 1e6:+,.1f}M VND", va="center", fontsize=9, color="#0F172A")
    ax1.set_xlim(0, t.records.max() * 1.9)
    ax1.set_title("Discrepancies by type", loc="left", fontsize=11)
    ax1.set_xlabel("Records")
    esc = a.variance_pct.abs() > RECON_TOLERANCE_PCT
    ax2.barh(a.advertiser_name, a.variance_pct, color=np.where(esc, "#DC2626", "#94A3B8"), height=0.65)
    for y, (v, e) in enumerate(zip(a.variance_pct, esc)):
        if e:
            ax2.text(v - 0.15 if v < 0 else v + 0.15, y, f"{v:+.1f}%", va="center",
                     ha="right" if v < 0 else "left", fontsize=8, color="#0F172A")
    for x in (-RECON_TOLERANCE_PCT, RECON_TOLERANCE_PCT):
        ax2.axvline(x, color="#DC2626", lw=1, ls="--")
    ax2.axvline(0, color="#0F172A", lw=0.8)
    ax2.set_xlim(min(-6, a.variance_pct.min() - 1.5), max(3.5, a.variance_pct.max() + 1.5))
    ax2.tick_params(axis="y", labelsize=8)
    ax2.set_xlabel("Unexplained variance vs platform revenue (%)")
    ax2.set_title(f"By advertiser - red = escalate (|variance| > {RECON_TOLERANCE_PCT:.0f}%)", loc="left", fontsize=11)
    for ax in (ax1, ax2):
        for sp in ("top", "right"):
            ax.spines[sp].set_visible(False)
    fig.suptitle(f"Advertiser reconciliation {month} (simulated advertiser reports)", x=0.01, ha="left",
                 fontsize=13, fontweight="bold", color="#0F172A")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--week", help="Monday of the week to report, e.g. 2026-09-21 (default: last complete week)")
    ap.add_argument("--month", help="Month to reconcile, e.g. 2026-09 (default: latest month in data)")
    a = ap.parse_args()

    os.makedirs(OUT_DIR, exist_ok=True)
    conv, clicks, offers, pubs = load()

    last_day = conv.conversion_time.max().normalize()
    week_start = pd.Timestamp(a.week) if a.week else last_day - pd.Timedelta(days=last_day.weekday() + 7)
    available = sorted(f[18:25] for f in os.listdir(ADV_DIR) if f.startswith("advertiser_report_"))
    month = a.month or available[-1]

    # weekly
    summary, detail, watch = build_weekly(conv, clicks, offers, week_start)
    wk_path = os.path.join(OUT_DIR, f"weekly_performance_{week_start:%Y-%m-%d}.xlsx")
    write_excel(wk_path, {"Summary": summary, "By Offer": detail, "Publisher Watchlist": watch},
                f"Weekly Performance {week_start:%d/%m} - {week_start + pd.Timedelta(days=6):%d/%m/%Y}")
    print(f"[OK] Weekly report  -> {os.path.relpath(wk_path, BASE)}")
    print(f"     Offers flagged: {(detail.flag != 'OK').sum()} / {len(detail)} | Publishers on watchlist: {len(watch)}")

    # monthly
    overview, by_adv, by_type, disc, payout = build_recon(conv, offers, pubs, month)
    mo_path = os.path.join(OUT_DIR, f"monthly_reconciliation_{month}.xlsx")
    write_excel(mo_path, {"Overview": overview, "By Advertiser": by_adv, "By Discrepancy Type": by_type,
                          "Discrepancy Detail": disc, "Publisher Payout": payout},
                f"Advertiser Reconciliation {month}")
    print(f"[OK] Reconciliation -> {os.path.relpath(mo_path, BASE)}")
    png = os.path.join(OUT_DIR, f"recon_{month}.png")
    plot_recon(by_adv, by_type, month, png)
    print(f"[OK] Chart          -> {os.path.relpath(png, BASE)}")
    for _, r in overview.iterrows():
        v = r.Value
        print(f"     {r.Item:<48}: {v:,.1f}" if isinstance(v, float) else f"     {r.Item:<48}: {v}")


if __name__ == "__main__":
    main()
